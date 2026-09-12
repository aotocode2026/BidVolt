"""浏览器内 Office 文件预览（issue #63，discussion #16 T02）。

策略：

- docx / doc / odt / rtf → 服务端 LibreOffice headless 转 PDF，前端用 PDF 查看器渲染；
- xlsx / xlsm / xltx（含 .xls 经 LibreOffice 转 xlsx）→ openpyxl 转单元格 JSON，前端表格渲染
  （Excel 直接转 PDF 受打印区域/分页影响，观感不可控）；
- pdf → 原件直出；
- 其他格式 → unsupported + 原因，由前端引导下载。

转换结果按内容寻址缓存到 `file_preview`：同一内容只转一次。转换是 CPU 密集型，
一律经线程池执行，避免阻塞事件循环。
"""

from __future__ import annotations

import asyncio
import io
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import date, datetime

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.file import FilePreview

DOC_EXTS = (".docx", ".doc", ".odt", ".rtf")
SHEET_EXTS = (".xlsx", ".xlsm", ".xltx")
LEGACY_SHEET_EXTS = (".xls",)
PDF_EXTS = (".pdf",)

# 与其它上传/产物通道保持一致的上限，避免超大文件拖垮转换进程
MAX_CONVERT_BYTES = 80 * 1024 * 1024
CONVERT_TIMEOUT_SECONDS = 240

# 表格预览载荷上限（超出部分由前端提示"仅显示前 N 行/列"）
SHEET_MAX_SHEETS = 12
SHEET_MAX_ROWS = 500
SHEET_MAX_COLS = 60
SHEET_MAX_CELLS = 80_000
CELL_MAX_CHARS = 500


def _normalize_ext(ext: str | None) -> str:
    value = str(ext or "").strip().lower()
    if value and not value.startswith("."):
        value = "." + value
    return value


def preview_kind_for_ext(ext: str | None) -> str:
    """预览类型：pdf（转 PDF 或原件）/ sheet（表格网格）/ unsupported。"""
    value = _normalize_ext(ext)
    if value in PDF_EXTS or value in DOC_EXTS:
        return "pdf"
    if value in SHEET_EXTS or value in LEGACY_SHEET_EXTS:
        return "sheet"
    return "unsupported"


@dataclass
class PreviewSource:
    """待预览文件的统一描述（上传材料 / 企业资料 / 成文产物都归一到此结构）。"""

    source_type: str  # file | artifact
    source_id: int
    enterprise_id: int
    filename: str
    ext: str
    mime: str
    version_key: str  # 文件 sha256 或 v<artifact 版本号>
    data: bytes

    @property
    def display_name(self) -> str:
        return str(self.filename or "").rsplit("/", 1)[-1]


def _soffice_convert_sync(data: bytes, filename: str, target_ext: str) -> bytes:
    """LibreOffice headless 转换（独立工作目录 + 独立 profile，天然支持并发）。"""
    suffix = os.path.splitext(str(filename or ""))[1].lower() or ".bin"
    work = tempfile.mkdtemp(prefix="bidvolt_preview_")
    try:
        src = os.path.join(work, "input" + suffix)
        with open(src, "wb") as fh:
            fh.write(data)
        env = dict(os.environ)
        env["HOME"] = work  # LibreOffice 需要可写 HOME，否则可能拒绝启动
        proc = subprocess.run(  # noqa: S603 固定命令，无 shell
            [
                "soffice",
                "--headless",
                "--norestore",
                "--nolockcheck",
                "--convert-to",
                target_ext,
                "--outdir",
                work,
                src,
                f"-env:UserInstallation=file://{work}/lo_profile",
            ],
            capture_output=True,
            timeout=CONVERT_TIMEOUT_SECONDS,
            env=env,
        )
        out = os.path.join(work, "input." + target_ext)
        if proc.returncode != 0 or not os.path.exists(out):
            detail = (proc.stderr or b"").decode("utf-8", "replace")[-300:]
            raise ValueError(
                f"文档转换失败（LibreOffice 返回 {proc.returncode}）：{detail or '无错误输出'}"
            )
        with open(out, "rb") as fh:
            return fh.read()
    except subprocess.TimeoutExpired as exc:
        raise ValueError(
            f"文档转换超时（{CONVERT_TIMEOUT_SECONDS}s）：文件可能过大或已损坏"
        ) from exc
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _pdf_page_count(pdf_bytes: bytes) -> int | None:
    try:
        import fitz  # PyMuPDF

        with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
            return int(doc.page_count)
    except Exception:  # noqa: BLE001 页数只是展示信息，取不到不影响预览
        return None


def _cell_value(value: object) -> object:
    """单元格值归一化为 JSON 安全类型。"""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value)[:CELL_MAX_CHARS]


def _sheet_payload_sync(data: bytes, filename: str) -> dict:
    """xlsx → 单元格网格 JSON。xls 先经 LibreOffice 转 xlsx 再解析。"""
    from openpyxl import load_workbook

    raw = data
    if _normalize_ext(os.path.splitext(str(filename or ""))[1]) in LEGACY_SHEET_EXTS:
        try:
            raw = _soffice_convert_sync(data, filename, "xlsx")
        except ValueError as exc:
            raise ValueError(
                "暂不支持 .xls 在线预览（服务端缺少表格转换组件），请另存为 .xlsx 后重新上传"
            ) from exc
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet_names = list(workbook.sheetnames)
        sheets: list[dict] = []
        total_cells = 0
        truncated = len(sheet_names) > SHEET_MAX_SHEETS
        for name in sheet_names[:SHEET_MAX_SHEETS]:
            worksheet = workbook[name]
            rows: list[list[object]] = []
            sheet_truncated = False
            for row in worksheet.iter_rows(values_only=True):
                if len(rows) >= SHEET_MAX_ROWS or total_cells >= SHEET_MAX_CELLS:
                    sheet_truncated = True
                    break
                if len(row) > SHEET_MAX_COLS:
                    sheet_truncated = True
                values = [_cell_value(v) for v in row[:SHEET_MAX_COLS]]
                while values and values[-1] in (None, ""):
                    values.pop()
                total_cells += len(values)
                rows.append(values)
            if sheet_truncated:
                truncated = True
            sheets.append(
                {
                    "name": name,
                    "rows": len(rows),
                    "cols": max((len(r) for r in rows), default=0),
                    "truncated": sheet_truncated,
                    "cells": rows,
                }
            )
        return {
            "sheets": sheets,
            "sheet_count": len(sheet_names),
            "truncated": truncated,
            "limits": {
                "max_sheets": SHEET_MAX_SHEETS,
                "max_rows": SHEET_MAX_ROWS,
                "max_cols": SHEET_MAX_COLS,
            },
        }
    finally:
        workbook.close()


async def _load_cached_pdf(session: AsyncSession, src: PreviewSource) -> FilePreview | None:
    return await session.scalar(
        select(FilePreview).where(
            FilePreview.enterprise_id == int(src.enterprise_id),
            FilePreview.source_type == src.source_type,
            FilePreview.source_key == src.version_key,
            FilePreview.kind == "pdf",
        )
    )


async def _ensure_pdf(session: AsyncSession, src: PreviewSource) -> tuple[FilePreview | None, bool]:
    """确保 PDF 缓存存在：命中直接返回，未命中转换并落库。返回 (行, 是否命中缓存)。"""
    cached = await _load_cached_pdf(session, src)
    if cached is not None:
        return cached, True
    if len(src.data) > MAX_CONVERT_BYTES:
        raise ValueError(
            f"文件超过 {MAX_CONVERT_BYTES // (1024 * 1024)}MB，暂不支持在线转换预览，请下载查看"
        )
    pdf_bytes = await asyncio.to_thread(
        _soffice_convert_sync, src.data, src.display_name, "pdf"
    )
    page_count = await asyncio.to_thread(_pdf_page_count, pdf_bytes)
    row = FilePreview(
        enterprise_id=int(src.enterprise_id),
        source_type=src.source_type,
        source_id=int(src.source_id),
        source_key=src.version_key,
        kind="pdf",
        mime="application/pdf",
        content=pdf_bytes,
        page_count=page_count,
        byte_size=len(pdf_bytes),
    )
    session.add(row)
    try:
        # 保存点隔离：并发请求同时转换同一文件时，唯一约束冲突只回滚本次 INSERT
        async with session.begin_nested():
            await session.flush()
    except IntegrityError:
        row = await _load_cached_pdf(session, src)
    await session.commit()
    # commit 会清掉事务级 RLS 上下文，后续若有其它库操作需重建
    from app.services.task_service import _set_rls_context  # noqa: PLC0415

    await _set_rls_context(session, src.enterprise_id)
    return row, False


async def build_manifest(session: AsyncSession, src: PreviewSource) -> dict:
    """预览清单：前端据此选择 PDF 查看器或表格渲染，unsupported 时引导下载。"""
    kind = preview_kind_for_ext(src.ext)
    base = {
        "kind": kind,
        "filename": src.display_name,
        "source_type": src.source_type,
        "source_id": int(src.source_id),
        "source_key": src.version_key,
        "mime": src.mime,
    }
    if kind == "unsupported":
        return {
            **base,
            "reason": f"暂不支持在浏览器预览 {_normalize_ext(src.ext) or '该'} 格式，请下载后查看",
        }
    if _normalize_ext(src.ext) in PDF_EXTS:
        return {**base, "mime": "application/pdf", "converted": False, "cached": True}
    if kind == "pdf":
        row, cached = await _ensure_pdf(session, src)
        if row is None:
            raise ValueError("预览转换失败，请稍后重试或下载查看")
        return {
            **base,
            "mime": "application/pdf",
            "converted": True,
            "cached": cached,
            "page_count": row.page_count,
            "byte_size": row.byte_size,
        }
    payload = await asyncio.to_thread(_sheet_payload_sync, src.data, src.display_name)
    return {**base, "converted": False, "cached": False, **payload}


async def get_pdf(session: AsyncSession, src: PreviewSource) -> tuple[bytes, str]:
    """PDF 字节（docx 走缓存转换；pdf 原件直出）。"""
    if _normalize_ext(src.ext) in PDF_EXTS:
        return src.data, "application/pdf"
    if preview_kind_for_ext(src.ext) != "pdf":
        raise ValueError("该文件不支持 PDF 预览")
    row, _cached = await _ensure_pdf(session, src)
    if row is None:
        raise ValueError("预览转换失败，请稍后重试或下载查看")
    return row.content or b"", "application/pdf"
