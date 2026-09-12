"""交付 docx 归一化回归：页码页脚 + 大纲级别噪声（issue #65 / #66）。

覆盖三类生产实测场景：
1. 整文件直写产物（`Document()` 从零新建）→ 包里没有页脚部件，归一化后必须有页码；
2. 底稿裁剪产物（202 场景）→ 含 PAGE 的页脚部件在包里但引用被删空，归一化要接回引用；
3. 图注/超深大纲级别（217 场景，306 条 `outlineLvl=8`）→ 必须清掉，标题级别保留。
"""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn

from app.services import docx_normalize


def _docx_bytes(paragraphs: list[tuple[str, int | None]] | list[str]) -> bytes:
    """构造 docx：元素为 (文本, outlineLvl) 或裸字符串（无大纲级别）。"""
    doc = Document()
    for item in paragraphs:
        text, lvl = (item, None) if isinstance(item, str) else item
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if lvl is not None:
            ppr = para._p.get_or_add_pPr()
            ol = ppr.makeelement(qn("w:outlineLvl"), {})
            ol.set(qn("w:val"), str(lvl))
            ppr.append(ol)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _read(content: bytes, name: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        return zf.read(name)


def _outline_levels(content: bytes) -> list[str]:
    xml = _read(content, "word/document.xml").decode("utf-8")
    return re.findall(r'<w:outlineLvl w:val="(\d+)"/>', xml)


# --------------------------------------------------------------------------- #
# 大纲级别噪声
# --------------------------------------------------------------------------- #


def test_strip_outline_noise_keeps_headings_removes_captions_and_deep_levels():
    raw = _docx_bytes(
        [
            ("第一章 项目理解", 0),  # 合法标题，保留
            ("1.1 现状分析", 5),  # 最深合法标题，保留
            ("图：营业执照扫描件-第1页", 8),  # 图注：清
            ("附件：审计报告", 3),  # 题注型段落：即使级别合法也要清
            ("超深标题", 8),  # outlineLvl ≥ 6：清
            "正文段落",  # 无级别，不动
        ]
    )
    out, removed = docx_normalize.strip_outline_noise(raw)
    assert removed == 3
    assert _outline_levels(out) == ["0", "5"]
    # 正文文字未被改动
    assert "第一章 项目理解" in _read(out, "word/document.xml").decode("utf-8")


def test_strip_outline_noise_idempotent():
    raw = _docx_bytes([("标题", 0), ("图：证据-第1页", 8)])
    once, removed = docx_normalize.strip_outline_noise(raw)
    assert removed == 1
    twice, removed_again = docx_normalize.strip_outline_noise(once)
    assert removed_again == 0
    assert twice == once


# --------------------------------------------------------------------------- #
# 页码页脚
# --------------------------------------------------------------------------- #


def test_ensure_page_footer_injects_template_style_footer():
    raw = _docx_bytes(["正文"])
    assert docx_normalize.audit_docx(raw)["has_page_footer"] is False

    out, stats = docx_normalize.ensure_page_footer(raw)
    assert stats["created_part"]
    info = docx_normalize.audit_docx(out)
    assert info["has_page_footer"] is True
    assert info["page_field_parts"] == 1
    assert info["sections_without_page_footer"] == 0
    # 模板口径：居中 + 单个 PAGE 域 + 9pt（sz=18）
    footer = _read(out, stats["created_part"]).decode("utf-8")
    assert 'w:jc w:val="center"' in footer
    assert "PAGE" in footer and "MERGEFORMAT" in footer
    assert 'w:sz w:val="18"' in footer
    assert "宋体" in footer
    # Content_Types / rels 同步登记
    ct = _read(out, "[Content_Types].xml").decode("utf-8")
    assert stats["created_part"].split("/")[-1] in ct
    rels = _read(out, "word/_rels/document.xml.rels").decode("utf-8")
    assert 'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer"' in rels
    # 节上挂了引用，且没有 start=0 的页码起始
    doc = _read(out, "word/document.xml").decode("utf-8")
    assert "<w:footerReference" in doc
    assert 'w:start="0"' not in doc


def test_ensure_page_footer_relinks_orphan_template_footer():
    """202 实测场景：含 PAGE 的页脚部件与 rels 都在，但节里的引用被删空 → 接回引用。"""
    raw = _docx_bytes(["商务偏差表"])
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        items = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    parts = dict(items)
    parts["word/footer3.xml"] = docx_normalize._footer_xml()
    rels = parts["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/footer" Target="footer3.xml"/></Relationships>',
        1,
    )
    parts["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    ct = parts["[Content_Types].xml"].decode("utf-8")
    parts["[Content_Types].xml"] = ct.replace(
        "</Types>",
        '<Override PartName="/word/footer3.xml" ContentType="application/vnd.'
        'openxmlformats-officedocument.wordprocessingml.footer+xml"/></Types>',
        1,
    ).encode("utf-8")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
        for name, blob in parts.items():
            zo.writestr(name, blob)
    orphan = buf.getvalue()

    assert docx_normalize.audit_docx(orphan)["has_page_footer"] is False
    out, stats = docx_normalize.ensure_page_footer(orphan)
    assert stats["created_part"] is None  # 复用既有部件，不新建
    assert stats["refs_added"] == 1
    assert docx_normalize.audit_docx(out)["has_page_footer"] is True


def test_ensure_page_footer_idempotent():
    raw = _docx_bytes(["正文"])
    once, first = docx_normalize.ensure_page_footer(raw)
    assert first["changed"] is True
    twice, second = docx_normalize.ensure_page_footer(once)
    assert second["changed"] is False
    assert twice == once


def test_ensure_page_footer_removes_start_zero_pgnum():
    """底稿封面节的 pgNumType start=0 会让首页显示「0」，交付件一律从 1 起。"""
    raw = _docx_bytes(["正文"])
    doc_xml = _read(raw, "word/document.xml").decode("utf-8")
    doc_xml = doc_xml.replace(
        "<w:sectPr", '<w:sectPr><w:pgNumType w:start="0"/>', 1
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as zf, zipfile.ZipFile(
        buf, "w", zipfile.ZIP_DEFLATED
    ) as zo:
        for name in zf.namelist():
            blob = doc_xml.encode("utf-8") if name == "word/document.xml" else zf.read(name)
            zo.writestr(name, blob)
    out, stats = docx_normalize.ensure_page_footer(buf.getvalue())
    assert stats["pg_num_start0_removed"] == 1
    assert 'w:start="0"' not in _read(out, "word/document.xml").decode("utf-8")


def test_normalize_docx_end_to_end():
    raw = _docx_bytes([("标题", 0), ("图：证据-第1页", 8), "正文"])
    out, stats = docx_normalize.normalize_docx(raw)
    assert stats["outline_noise_removed"] == 1
    assert stats["created_part"]
    info = docx_normalize.audit_docx(out)
    assert info["has_page_footer"] is True
    assert info["outline_noise"] == 0
    # 非 docx 原样返回
    same, none_stats = docx_normalize.normalize_docx(b"not-a-zip")
    assert same == b"not-a-zip"
    assert none_stats == {"is_docx": False}
