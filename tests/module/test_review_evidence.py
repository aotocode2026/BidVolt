"""评分输入构造回归（issue #70）：分块索引 / 要素清单取文 / 覆盖情况 / 编制记录索引。"""

from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn

from app.services import review_evidence as ev


def _docx_bytes(items: list[tuple[str, int | None]], table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for text, level in items:
        p = doc.add_paragraph()
        p.add_run(text)
        if level is not None:
            ppr = p._p.get_or_add_pPr()
            ol = ppr.makeelement(qn("w:outlineLvl"), {})
            ol.set(qn("w:val"), str(level - 1))
            ppr.append(ol)
    if table:
        tb = doc.add_table(rows=len(table), cols=len(table[0]))
        for ri, row in enumerate(table):
            for ci, cell in enumerate(row):
                tb.cell(ri, ci).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_chunk_docx_splits_by_heading_and_keeps_path():
    data = _docx_bytes(
        [
            ("一、项目理解", 1),
            ("本节说明项目背景。", None),
            ("3.1 项目背景与政策形势理解", 2),
            ("虚拟电厂相关政策……", None),
            ("三、总体方案", 1),
            ("总体架构分三层……", None),
        ]
    )
    counter = [0]
    chunks = ev.chunk_docx(data, 42, "技术文件/（二）专项响应文件.docx", counter)
    paths = [c.heading_path for c in chunks]
    assert any("一、项目理解" in p for p in paths)
    assert any("3.1 项目背景与政策形势理解" in p for p in paths)
    assert any("三、总体方案" in p for p in paths)
    # 标题路径保留层级：二级标题的路径包含一级标题
    deep = [c for c in chunks if "3.1" in c.heading_path]
    assert deep and "一、项目理解" in deep[0].heading_path
    assert all(c.artifact_id == 42 for c in chunks)
    assert all(c.file_name.endswith("专项响应文件.docx") for c in chunks)


def test_chunk_docx_splits_long_section_into_windows():
    long_text = "长段落内容" * 4000  # 20,000 字
    data = _docx_bytes([("5.8 商务评分支撑材料", 2), (long_text, None)])
    counter = [0]
    chunks = ev.chunk_docx(data, 7, "商务文件/（四）补充文件.docx", counter)
    assert len(chunks) >= 3  # 8000 字一块 → 至少 3 块
    assert any("（续" in c.label for c in chunks)


def _chunks() -> list[ev.Chunk]:
    return [
        ev.Chunk(1, 1, "技术文件/专项响应文件.docx", "5.8.2 创新激励机制", "创新激励" * 50),
        ev.Chunk(2, 1, "技术文件/专项响应文件.docx", "5.8.3 研发团队规模", "高级技师" * 30),
        ev.Chunk(3, 2, "商务文件/补充文件.docx", "七、进度计划", "里程碑" * 40),
    ]


def test_select_chunks_scores_by_keys_and_respects_budget():
    checklist = [
        {"element": "创新激励制度", "keys": ["创新激励", "研发奖励"]},
        {"element": "高职称人数", "keys": ["高级技师"]},
    ]
    selected, scope = ev.select_chunks(checklist, _chunks(), budget=1000)
    assert [c.index for c in selected] == [1, 2]
    assert scope["chunks_matched"] == 2
    assert scope["chars_provided"] <= 1000
    assert scope["fallback"] is False


def test_select_chunks_zero_hit_triggers_fallback_signal():
    checklist = [{"element": "采购人绩效评价", "keys": ["绩效评价", "评价等级"]}]
    selected, scope = ev.select_chunks(checklist, _chunks())
    assert selected == []
    assert scope["chunks_matched"] == 0  # 调用方据此启用兜底选块


def test_select_chunks_extra_keys_work_without_checklist():
    """要素清单缺失时，规则原文词仍能取到正文（issue #70 首次上线回归的修复）。"""
    selected, scope = ev.select_chunks([], _chunks(), extra_keys=["里程碑"])
    assert [c.index for c in selected] == [3]
    assert scope["chunks_matched"] == 1


def test_keys_from_rule_text_extracts_terms():
    keys = ev.keys_from_rule_text("服务方案、管理组织、设备设施配置的先进性、创新性……优27-30分")
    assert "服务方案" in keys
    assert "管理组织" in keys
    assert len(keys) <= 12


def test_fallback_chunks_by_category_prefers_matching_files():
    chunks = [
        ev.Chunk(1, 1, "技术文件/（二）专项响应文件.docx", "四、工作规划描述", "方案" * 100),
        ev.Chunk(2, 2, "商务文件/（四）补充文件.docx", "八、科研创新", "激励" * 100),
        ev.Chunk(3, 3, "内部管理文件/编制逻辑与评分响应记录.docx", "三、装订矩阵", "索引" * 100),
    ]
    tech = ev.fallback_chunks_by_category("技术", chunks, budget=1000)
    assert [c.index for c in tech] == [1, 3]
    biz = ev.fallback_chunks_by_category("商务", chunks, budget=1000)
    assert [c.index for c in biz] == [2, 3]


def test_build_checklists_retries_missing_rules(monkeypatch):
    """批内漏掉的规则必须逐条补生成（第一次上线：批量调用解析失败 → 全部清单为空）。"""
    import asyncio
    import json as _json

    from app.services.llm import LLMClient

    calls: list[str] = []

    async def fake_chat(self, system, user):  # noqa: ANN001
        calls.append(user)
        if "[rule_index=0]" in user and "[rule_index=1]" in user:
            return "not-a-json"  # 整批解析失败
        return _json.dumps(
            {"checklists": [{"rule_index": int(user.split("rule_index=")[1][:1]),
                             "elements": [{"element": "要素", "keys": ["关键词"]}]}]},
            ensure_ascii=False,
        )

    monkeypatch.setattr(LLMClient, "chat", fake_chat)
    rules = [{"content": "规则A", "category": "技术", "weight": 10},
             {"content": "规则B", "category": "商务", "weight": 5}]
    out = asyncio.run(ev.build_checklists(rules, batch_size=5))
    assert set(out) == {0, 1}  # 批量失败后逐条补齐
    assert len(calls) == 3


def test_select_by_index_for_fallback():
    selected = ev.select_by_index(_chunks(), [3, 1], budget=600)
    assert [c.index for c in selected] == [1, 3]


def test_chunk_map_respects_budget():
    chunks = _chunks()
    text = ev.chunk_map(chunks, budget=60)
    assert len(text) <= 200  # 截断提示会略超，但不会把全文塞进来
    assert "块" in text


def test_build_editorial_hints_extracts_matrix_and_risk():
    data = _docx_bytes(
        [
            ("三、评分装订矩阵与技术/商务评分响应", None),
            ("九、未闭环/风险提示", None),
            ("a. 批次已成交（复盘口径）；b. 少数栏目留白。", None),
        ],
        table=[
            ["评分点", "分值", "证明材料", "响应位置", "预计得分"],
            ["技术2a 项目负责人", "10", "简历表+身份证", "团队章节", "视证据"],
        ],
    )
    hints = ev.build_editorial_hints(data)
    assert hints["matrix"] and hints["matrix"][0][0] == "评分点"
    assert "留白" in hints["risk"]

    block = ev.editorial_block(hints, "项目负责人专业、工作年限、经验评价", ["项目负责人"])
    assert "预计得分" in block
    assert "严禁作为得分依据" in block or "风险提示" in block


def test_is_editorial_matches_internal_record():
    assert ev.is_editorial("内部管理文件/编制逻辑与评分响应记录.docx")
    assert not ev.is_editorial("技术文件/（二）专项响应文件.docx")


# --------------------------------------------------------------------------- #
# 评审提示词装配（issue #70）：要素清单 / 目录 / 定向取文 / 编制方自述必须都进 prompt
# --------------------------------------------------------------------------- #


def test_score_rule_prompt_includes_checklist_map_chunks_and_editorial(monkeypatch):
    import asyncio
    import json as _json

    from app.services import review_service
    from app.services.llm import LLMClient

    captured: dict = {}

    async def fake_chat(self, system, user):  # noqa: ANN001
        captured["system"] = system
        captured["user"] = user
        return _json.dumps(
            {
                "verdict": "partial",
                "got": 6,
                "element_results": [
                    {"element": "项目负责人学历", "result": "hit", "quote": "学历：本科"}
                ],
                "evidence_scope": {"chunks_provided": 1, "chunks_total": 3, "scope_limited": False},
            },
            ensure_ascii=False,
        )

    monkeypatch.setattr(LLMClient, "chat", fake_chat)

    rule = {"content": "项目负责人专业、工作年限、经验评价", "weight": 10, "category": "技术"}
    context = {
        "checklist": [{"element": "项目负责人学历", "keys": ["学历"], "criteria": "本科以上"}],
        "scope": {"chunks_total": 3, "chars_total": 9000, "chunks_provided": 1, "chars_provided": 3000},
        "map": "[块1] 技术文件 › 二、项目团队情况（3000字）：简历表",
        "chunks_text": "【块1｜技术文件 › 二、项目团队情况】\n项目负责人简历表：陈千里",
        "editorial": "技术2a 项目负责人 | 10 | 简历表+身份证 | 团队章节 | 视证据",
    }
    parsed = asyncio.run(review_service._score_rule_with_llm(rule, context))
    user = captured["user"]
    assert "核验要素清单" in user and "项目负责人学历" in user
    assert "交付件分块目录" in user and "[块1]" in user
    assert "相关章节全文" in user and "陈千里" in user
    assert "编制方自述" in user and "严禁作为得分依据" in user
    assert parsed["verdict"] == "partial"
    assert parsed["element_results"][0]["result"] == "hit"
